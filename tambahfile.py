import getpass
import hashlib
import json
import secrets
from pathlib import Path

from openpyxl import Workbook, load_workbook


FILE_JSON = Path(__file__).with_name("akun.json")
FILE_EXCEL = Path(__file__).with_name("akun.xlsx")
ITERASI_HASH = 600_000
HEADER_EXCEL_TERBACA = ("username", "password")


def _buat_hash_password(password, salt):
    return hashlib.pbkdf2_hmac(
        "sha256", password.encode("utf-8"), bytes.fromhex(salt), ITERASI_HASH
    ).hex()


def _simpan_excel(akun, username_baru, password_baru):
    password_tersimpan = {}
    if FILE_EXCEL.exists():
        workbook_lama = load_workbook(FILE_EXCEL, read_only=True, data_only=True)
        try:
            worksheet_lama = workbook_lama.active
            baris = worksheet_lama.iter_rows(values_only=True)
            header = next(baris, ())
            if "username" in header and "password" in header:
                kolom_username = header.index("username")
                kolom_password = header.index("password")
                for data in baris:
                    if data[kolom_username] is not None:
                        password_tersimpan[str(data[kolom_username]).casefold()] = (
                            data[kolom_password] or ""
                        )
        finally:
            workbook_lama.close()

    password_tersimpan[username_baru.casefold()] = password_baru
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Akun"
    worksheet.append(HEADER_EXCEL_TERBACA)

    for data in akun:
        username = data["username"]
        worksheet.append((username, password_tersimpan.get(username.casefold(), "")))
        for cell in worksheet[worksheet.max_row]:
            cell.data_type = "s"

    workbook.save(FILE_EXCEL)


def tambah_akun():
    """Tambahkan akun ke file JSON dan ekspor seluruh data ke Excel."""
    if FILE_JSON.exists():
        with FILE_JSON.open("r", encoding="utf-8") as file:
            akun = json.load(file)
        if not isinstance(akun, list):
            raise ValueError("Format data akun.json harus berupa daftar.")
    else:
        akun = []

    username = input("Masukkan username: ").strip()
    if not username:
        print("Username tidak boleh kosong.")
        return

    if any(data.get("username", "").casefold() == username.casefold() for data in akun):
        print("Username tersebut sudah terdaftar.")
        return

    password = getpass.getpass("Masukkan password: ")
    if not password:
        print("Password tidak boleh kosong.")
        return
    if password != getpass.getpass("Ulangi password: "):
        print("Password tidak cocok. Data tidak disimpan.")
        return

    salt = secrets.token_hex(16)
    akun.append(
        {
            "username": username,
            "password_hash": _buat_hash_password(password, salt),
            "salt": salt,
        }
    )

    with FILE_JSON.open("w", encoding="utf-8") as file:
        json.dump(akun, file, ensure_ascii=False, indent=2)

    _simpan_excel(akun, username, password)
    print("Data akun berhasil disimpan ke akun.json dan akun.xlsx.")
    print("Catatan: password di akun.xlsx tersimpan sebagai teks biasa.")